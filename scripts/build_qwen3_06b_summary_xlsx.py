#!/usr/bin/env python3
"""Aggregate the Qwen3-0.6B campaign results into a single xlsx workbook."""
import glob, json, os
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

ROOT = Path("results/eval/qwen3-0.6b")
OUT = ROOT / "SUMMARY-qwen3-0.6b.xlsx"

LM_TASKS = ["hellaswag", "arc_challenge", "mmlu", "truthfulqa_mc2", "winogrande", "gsm8k"]
LM_LABELS = {"hellaswag":"HellaSwag", "arc_challenge":"ARC-c", "mmlu":"MMLU",
             "truthfulqa_mc2":"TruthfulQA", "winogrande":"WinoGrande", "gsm8k":"GSM8K"}
LM_METRIC = {"hellaswag":"acc_norm,none", "arc_challenge":"acc_norm,none", "mmlu":"acc,none",
             "truthfulqa_mc2":"acc,none", "winogrande":"acc,none", "gsm8k":"exact_match,strict-match"}
MODEL_ORDER = ["teacher_sft", "teacher_dpo", "student_sft",
               "dpo", "dckd", "vanilakd", "adpa", "tvkd",
               "riskkd", "pfw_tail"]
GROUP = {**dict.fromkeys(["teacher_sft","teacher_dpo"], "Teacher (1.7B)"),
         "student_sft": "Student init (0.6B)",
         **dict.fromkeys(["dpo","dckd","vanilakd","adpa","tvkd"], "Baseline (0.6B)"),
         **dict.fromkeys(["riskkd","pfw_tail"], "Ours (0.6B)")}

HDR_FONT = Font(bold=True, color="FFFFFF")
HDR_FILL = PatternFill("solid", fgColor="305496")
OURS_FILL = PatternFill("solid", fgColor="E2EFDA")
TEACHER_FILL = PatternFill("solid", fgColor="FFF2CC")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CENTER = Alignment(horizontal="center", vertical="center")
LEFT = Alignment(horizontal="left", vertical="center")
BOLD = Font(bold=True)

def style_header(ws, row, ncols):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = HDR_FONT; cell.fill = HDR_FILL
        cell.alignment = CENTER; cell.border = BORDER

def row_fill_for(name):
    if name in ("teacher_sft","teacher_dpo"): return TEACHER_FILL
    if name in ("riskkd","pfw_tail"): return OURS_FILL
    return None

def load_lm():
    out = {}
    for jf in glob.glob(str(ROOT / "lm_eval/*/*/results*.json")):
        name = jf.split("/")[4]
        d = json.load(open(jf))
        res = d["results"]
        row = {t: res.get(t, {}).get(LM_METRIC[t]) for t in LM_TASKS}
        row["gsm8k_flex"] = res.get("gsm8k", {}).get("exact_match,flexible-extract")
        out[name] = row
    return out

def load_mt():
    rows = []
    for jf in sorted(glob.glob(str(ROOT / "mt_bench_pairwise/gpt-5.4-mini/*.json"))):
        d = json.load(open(jf))
        a = d["anchor"].replace("qwen3-0.6b-", "")
        b = d["baseline"].replace("qwen3-0.6b-", "")
        ag = d["aggregate"]
        rows.append({
            "anchor": a, "baseline": b, "judge": d["judge"],
            "n_questions": d["n_questions"],
            "wins": ag["wins"], "ties": ag["ties"], "losses": ag["losses"], "errors": ag.get("errors", 0),
            "win_rate": ag["win_rate"],
            "adj_win_rate": (ag["wins"] + 0.5 * ag["ties"]) / max(1, ag["wins"] + ag["ties"] + ag["losses"]),
            "win_rate_turn1": ag.get("win_rate_turn1"),
            "win_rate_turn2": ag.get("win_rate_turn2"),
        })
    return rows

def sheet_lm(wb, lm):
    ws = wb.create_sheet("lm-eval")
    hdr = ["Group", "Model"] + [LM_LABELS[t] for t in LM_TASKS] + ["GSM8K (flex)", "AVG (6 tasks)"]
    ws.append(hdr); style_header(ws, 1, len(hdr))
    for m in MODEL_ORDER:
        r = lm.get(m, {})
        vals = [r.get(t) for t in LM_TASKS]
        nums = [v for v in vals if v is not None]
        avg = (sum(nums) / len(nums)) if nums else None
        ws.append([GROUP[m], m] + vals + [r.get("gsm8k_flex"), avg])
        rownum = ws.max_row
        fill = row_fill_for(m)
        for c in range(1, len(hdr) + 1):
            cell = ws.cell(row=rownum, column=c)
            cell.border = BORDER
            if fill: cell.fill = fill
            cell.alignment = LEFT if c < 3 else CENTER
            if c >= 3: cell.number_format = "0.00%"
            if c == len(hdr): cell.font = BOLD
    widths = [22, 14] + [12] * len(LM_TASKS) + [13, 14]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "C2"

def sheet_mt(wb, mt):
    ws = wb.create_sheet("mt-bench")
    hdr = ["Anchor (Ours)", "Baseline", "Judge", "N (Q x 2 turns x 2 pos)",
           "Wins", "Ties", "Losses", "Errors",
           "Win rate", "Adj win rate (W+0.5T)/(W+T+L)",
           "Win rate (turn 1)", "Win rate (turn 2)"]
    ws.append(hdr); style_header(ws, 1, len(hdr))
    for r in mt:
        n_total = r["wins"] + r["ties"] + r["losses"] + r["errors"]
        n_label = f"{r['n_questions']}x2x2 = {n_total}"
        ws.append([r["anchor"], r["baseline"], r["judge"], n_label,
                   r["wins"], r["ties"], r["losses"], r["errors"],
                   r["win_rate"], r["adj_win_rate"],
                   r["win_rate_turn1"], r["win_rate_turn2"]])
        rownum = ws.max_row
        for c in range(1, len(hdr) + 1):
            cell = ws.cell(row=rownum, column=c)
            cell.border = BORDER
            cell.alignment = CENTER if c >= 3 else LEFT
            if c >= 9: cell.number_format = "0.0%"
    widths = [14, 12, 14, 22, 7, 7, 8, 8, 11, 24, 16, 16]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "C2"

def sheet_summary(wb, lm, mt):
    ws = wb.create_sheet("summary", 0)
    ws["A1"] = "Qwen3-0.6B campaign - eval summary"
    ws["A1"].font = Font(bold=True, size=14)
    ws.merge_cells("A1:H1")
    ws["A2"] = ("Teacher = Qwen3-1.7B (SFT on Deita, then DPO on UltraFeedback). "
                "Student = Qwen3-0.6B-Base SFT'd on Deita as init for all distillation/DPO methods. "
                "lm-eval: vLLM TP=1, 6 tasks. MT-Bench: 80 Qs x 2 turns x 2 positions (g1/g2) = 320 judgments per pair, "
                "judge = gpt-5.4-mini.")
    ws["A2"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells("A2:H3")
    ws.row_dimensions[2].height = 34
    ws.row_dimensions[3].height = 22

    ws["A5"] = "lm-eval scores"
    ws["A5"].font = HDR_FONT; ws["A5"].fill = HDR_FILL
    hdr = ["Model", "AVG", "HellaSwag", "ARC-c", "MMLU", "TruthfulQA", "WinoGrande", "GSM8K"]
    row0 = 6
    for i, h in enumerate(hdr, 1):
        c = ws.cell(row=row0, column=i, value=h)
        c.font = HDR_FONT; c.fill = HDR_FILL; c.alignment = CENTER; c.border = BORDER
    for m in MODEL_ORDER:
        r = lm.get(m, {})
        vals = [r.get(t) for t in LM_TASKS]
        nums = [v for v in vals if v is not None]
        avg = sum(nums) / len(nums) if nums else None
        ws.append([m, avg] + vals)
        rownum = ws.max_row
        fill = row_fill_for(m)
        for c in range(1, len(hdr) + 1):
            cell = ws.cell(row=rownum, column=c)
            cell.border = BORDER
            if fill: cell.fill = fill
            cell.alignment = LEFT if c == 1 else CENTER
            if c >= 2: cell.number_format = "0.00%"
            if c == 2: cell.font = BOLD

    base_row = ws.max_row + 3
    ws.cell(row=base_row, column=1, value="MT-Bench win rate (Ours vs each baseline)").font = HDR_FONT
    ws.cell(row=base_row, column=1).fill = HDR_FILL
    anchors = sorted({r["anchor"] for r in mt})
    baselines = sorted({r["baseline"] for r in mt})
    hdr2 = ["Anchor \\ Baseline"] + baselines + ["mean"]
    for i, h in enumerate(hdr2, 1):
        c = ws.cell(row=base_row + 1, column=i, value=h)
        c.font = HDR_FONT; c.fill = HDR_FILL; c.alignment = CENTER; c.border = BORDER
    lookup = {(r["anchor"], r["baseline"]): r["win_rate"] for r in mt}
    for ai, a in enumerate(anchors):
        rownum = base_row + 2 + ai
        ws.cell(row=rownum, column=1, value=a).font = BOLD
        wins = []
        for bi, b in enumerate(baselines):
            v = lookup.get((a, b))
            if v is not None: wins.append(v)
            c = ws.cell(row=rownum, column=2 + bi, value=v)
            c.number_format = "0.0%"; c.alignment = CENTER; c.border = BORDER
            if v is not None and v > 0.5: c.fill = OURS_FILL
        mean = sum(wins) / len(wins) if wins else None
        c = ws.cell(row=rownum, column=2 + len(baselines), value=mean)
        c.number_format = "0.0%"; c.alignment = CENTER; c.font = BOLD; c.border = BORDER
        if mean is not None and mean > 0.5: c.fill = OURS_FILL

    note_row = ws.max_row + 3
    ws.cell(row=note_row, column=1,
            value="AlpacaEval: not run (datasets>=3.0 incompat with alpaca_eval script loader)."
            ).font = Font(italic=True, color="7F7F7F")
    ws.merge_cells(start_row=note_row, start_column=1, end_row=note_row, end_column=8)

    widths = [22, 11, 12, 11, 11, 12, 12, 11]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

def main():
    lm = load_lm()
    mt = load_mt()
    wb = Workbook(); wb.remove(wb.active)
    sheet_summary(wb, lm, mt)
    sheet_lm(wb, lm)
    sheet_mt(wb, mt)
    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    print(f"Wrote {OUT}")
    print(f"  lm-eval models: {len(lm)}    MT-Bench pairs: {len(mt)}")

if __name__ == "__main__":
    main()
