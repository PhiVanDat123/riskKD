"""Append pairwise MT-Bench sheets to the existing SUMMARY-mtbench-alpaca.xlsx.

Reads results/eval/mtbench-pair-v2/<judge>/<anchor>__vs__<baseline>.json and
adds three sheets:
  1. Pairwise Qwen       - 3 Ours x 6 baselines matrix (WR cells, color-coded)
  2. Pairwise Llama      - 3 Ours x 3 baselines matrix
  3. Pairwise Details    - per-pair W/L/T/n_total/win_rate
"""
import json
from pathlib import Path

from openpyxl import load_workbook, Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

REPO = Path("/Users/kienvu/Desktop/research-papers/riskKD")
XLSX = REPO / "results/eval/SUMMARY-mtbench-alpaca.xlsx"
PAIR_DIR = REPO / "results/eval/mtbench-pair-v2"
JUDGE = "gpt-5.4-mini"  # primary judge for the matrix

# Anchor + baseline orderings
ANCHORS_QWEN = ["qwen3-1.7b-tailriskKD", "qwen3-1.7b-pfw-tail", "qwen3-1.7b-riskKD"]
QWEN_BASELINES = [
    "qwen3-1.7b-ultrafeedback-dpo",
    "qwen3-1.7b-ultrafeedback-wpo",
    "qwen3-1.7b-ultrafeedback-radpo",
    "qwen3-1.7b-ultrafeedback-dckd",
    "qwen3-1.7b-ultrafeedback-adpa",
    "qwen3-1.7b-ultrafeedback-tvkd",
]
ANCHORS_LLAMA = ["llama3.2-1b-tailriskKD", "llama3.2-1b-pfw-tail", "llama-3.2-1b-riskkd-tokenwt-epoch1"]
LLAMA_BASELINES = [
    "llama-3.2-1b-dckd-ultrafeedback",
    "llama-3.2-1b-adpa-ultrafeedback",
    "llama-3.2-1b-tvkd-ultrafeedback",
]


def short_anchor(s):
    return (s.replace("qwen3-1.7b-", "")
             .replace("llama3.2-1b-", "")
             .replace("llama-3.2-1b-", "")
             .replace("-tokenwt-epoch1", " (tokenwt)"))


def short_baseline(s):
    return (s.replace("qwen3-1.7b-ultrafeedback-", "")
             .replace("qwen3-1.7b-", "")
             .replace("llama-3.2-1b-", "")
             .replace("-ultrafeedback", ""))


def load_pair(anchor, baseline, judge=JUDGE):
    path = PAIR_DIR / judge / f"{anchor}__vs__{baseline}.json"
    if not path.exists():
        return None
    return json.load(open(path))


HEADER_FILL = PatternFill(start_color="305496", end_color="305496", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True)
WIN_FILL = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
LOSS_FILL = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
TIE_FILL = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")
BORDER = Border(*[Side(style="thin", color="999999")] * 4)


def color_cell(cell, wr):
    if wr is None:
        return
    if wr > 0.55:
        cell.fill = WIN_FILL
    elif wr < 0.45:
        cell.fill = LOSS_FILL
    else:
        cell.fill = TIE_FILL


def write_header_row(ws, row, values):
    for i, v in enumerate(values, 1):
        c = ws.cell(row=row, column=i, value=v)
        c.fill = HEADER_FILL
        c.font = HEADER_FONT
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = BORDER


def auto_width(ws, ncols, max_width=30):
    for c in range(1, ncols + 1):
        col_letter = get_column_letter(c)
        max_len = 0
        for cell in ws[col_letter]:
            v = cell.value
            if v is None:
                continue
            l = len(str(v))
            if l > max_len:
                max_len = l
        ws.column_dimensions[col_letter].width = min(max_len + 3, max_width)


def add_matrix_sheet(wb, sheet_name, title, anchors, baselines):
    if sheet_name in wb.sheetnames:
        del wb[sheet_name]
    ws = wb.create_sheet(sheet_name)
    ws.cell(row=1, column=1, value=title).font = Font(bold=True, size=14)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(baselines) + 1)
    ws.cell(row=2, column=1,
            value=f"Each cell = WR = (wins + 0.5*ties) / total over 80Q x 2 turns x 2 swaps. Judge: {JUDGE}.").font = Font(italic=True, size=10)
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(baselines) + 1)
    ws.cell(row=3, column=1, value="GREEN = Ours WIN (WR>0.55) | YELLOW = TIE | RED = Baseline win (WR<0.45)").font = Font(italic=True, size=10)
    ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=len(baselines) + 1)

    write_header_row(ws, 5, ["Anchor (Ours)"] + [f"vs {short_baseline(b)}" for b in baselines])
    for row_i, anchor in enumerate(anchors, start=6):
        ws.cell(row=row_i, column=1, value=short_anchor(anchor)).font = Font(bold=True)
        ws.cell(row=row_i, column=1).border = BORDER
        for col_i, baseline in enumerate(baselines, start=2):
            d = load_pair(anchor, baseline)
            wr = d["aggregate"]["win_rate"] if d else None
            c = ws.cell(row=row_i, column=col_i, value=wr)
            c.border = BORDER
            c.alignment = Alignment(horizontal="center")
            if wr is not None:
                c.number_format = "0.000"
                color_cell(c, wr)
    auto_width(ws, len(baselines) + 1)


def add_details_sheet(wb):
    sheet_name = "Pairwise Details"
    if sheet_name in wb.sheetnames:
        del wb[sheet_name]
    ws = wb.create_sheet(sheet_name)
    write_header_row(ws, 1, ["Family", "Anchor (Ours)", "Baseline", "Wins", "Losses", "Ties",
                              "Errors", "n_total", "Win Rate", "WR Turn 1", "WR Turn 2", "Tokens In", "Tokens Out"])
    row = 2
    for family, anchors, baselines in [
        ("qwen3-1.7b", ANCHORS_QWEN, QWEN_BASELINES),
        ("llama-3.2-1b", ANCHORS_LLAMA, LLAMA_BASELINES),
    ]:
        for anchor in anchors:
            for baseline in baselines:
                d = load_pair(anchor, baseline)
                if d is None:
                    continue
                a = d["aggregate"]
                ws.cell(row=row, column=1, value=family)
                ws.cell(row=row, column=2, value=short_anchor(anchor))
                ws.cell(row=row, column=3, value=short_baseline(baseline))
                ws.cell(row=row, column=4, value=a.get("wins"))
                ws.cell(row=row, column=5, value=a.get("losses"))
                ws.cell(row=row, column=6, value=a.get("ties"))
                ws.cell(row=row, column=7, value=a.get("errors"))
                ws.cell(row=row, column=8, value=a.get("n_total"))
                wr_cell = ws.cell(row=row, column=9, value=a.get("win_rate"))
                wr_cell.number_format = "0.000"
                color_cell(wr_cell, a.get("win_rate"))
                ws.cell(row=row, column=10, value=a.get("win_rate_turn1")).number_format = "0.000"
                ws.cell(row=row, column=11, value=a.get("win_rate_turn2")).number_format = "0.000"
                ws.cell(row=row, column=12, value=d.get("total_input_tokens"))
                ws.cell(row=row, column=13, value=d.get("total_output_tokens"))
                for c in range(1, 14):
                    ws.cell(row=row, column=c).border = BORDER
                    if c >= 4:
                        ws.cell(row=row, column=c).alignment = Alignment(horizontal="center")
                row += 1
    auto_width(ws, 13)


def main():
    if not XLSX.exists():
        print(f"ERROR: {XLSX} not found. Run build_eval_xlsx.py first.")
        return 1
    wb = load_workbook(XLSX)
    add_matrix_sheet(wb, "Pairwise Qwen", "Pairwise MT-Bench - Qwen 1.7B Ours vs baselines",
                     ANCHORS_QWEN, QWEN_BASELINES)
    add_matrix_sheet(wb, "Pairwise Llama", "Pairwise MT-Bench - Llama 3.2-1B Ours vs baselines",
                     ANCHORS_LLAMA, LLAMA_BASELINES)
    add_details_sheet(wb)
    wb.save(XLSX)
    print(f"Updated {XLSX} with pairwise sheets")
    print(f"  Sheets: {wb.sheetnames}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
